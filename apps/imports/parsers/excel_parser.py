from datetime import date, datetime
from decimal import Decimal

from django.conf import settings
from openpyxl import load_workbook

from apps.imports.exceptions import ImportValidationError


class ExcelParser:
    OUTPUT_FIELDS = (
        "Registro",
        "Co",
        "Serie",
        "Descripcion",
        "Comienzo",
        "Vencimiento",
        "Cliente",
        "Usuario",
        "Fechareg",
        "EstadoArtA",
        "GrupoKarde",
        "Cierre",
    )
    HEADER_ALIASES = {
        "registro": "Registro",
        "co": "Co",
        "codigo": "Co",
        "serie": "Serie",
        "descripcion": "Descripcion",
        "comienzo": "Comienzo",
        "vencimiento": "Vencimiento",
        "cliente": "Cliente",
        "usuario": "Usuario",
        "fechareg": "Fechareg",
        "estadoarta": "EstadoArtA",
        "estadoartalqui": "EstadoArtA",
        "grupokarde": "GrupoKarde",
        "grupokardex": "GrupoKarde",
        "cierre": "Cierre",
    }

    def parse(self, uploaded_file):
        self._validate_file(uploaded_file)
        try:
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        except Exception as exc:
            raise ImportValidationError("El archivo Excel no es válido.", {"file": "corrupt"}) from exc

        sheet_name = settings.IMPORT_EXCEL_SHEET_NAME
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ImportValidationError("La hoja solicitada no existe.", {"sheet": sheet_name})
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook[workbook.sheetnames[0]]

        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            return []

        headers = [str(value).strip() if value is not None else "" for value in raw_headers]

        records = []
        for row in rows:
            values = list(row)
            values.extend([None] * (len(headers) - len(values)))
            records.append(self._build_record(headers, values))

        return records

    @classmethod
    def _build_record(cls, headers, values):
        record = dict.fromkeys(cls.OUTPUT_FIELDS)
        for header, value in zip(headers, values):
            field = cls.HEADER_ALIASES.get(header.casefold())
            if field:
                record[field] = cls._normalize(value, field)
        return record

    @staticmethod
    def _validate_file(uploaded_file):
        if not uploaded_file:
            raise ImportValidationError("Debe enviar un archivo.", {"file": "required"})
        if not uploaded_file.name.lower().endswith(".xlsx"):
            raise ImportValidationError("La extensión debe ser .xlsx.", {"file": "invalid_extension"})
        if uploaded_file.size > settings.IMPORT_MAX_FILE_SIZE_BYTES:
            raise ImportValidationError("El archivo excede el tamaño permitido.", {"file": "too_large"})

    @staticmethod
    def _normalize(value, field=None):
        if isinstance(value, (datetime, date)):
            return value.strftime("%d/%m/%Y %H:%M") if isinstance(value, datetime) else value.strftime("%d/%m/%Y")
        if isinstance(value, Decimal):
            value = float(value)
        if field == "Registro" and isinstance(value, float) and value.is_integer():
            return int(value)
        if field == "Cierre" and isinstance(value, bool):
            return "VERDADERO" if value else "FALSO"
        return value